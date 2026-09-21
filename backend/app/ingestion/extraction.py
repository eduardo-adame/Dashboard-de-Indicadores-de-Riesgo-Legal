"""Extracción inicial sin reglas de Validation ni OCR."""
from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.ingestion.csv_policy import CsvStructureError, determine_delimiter
from app.ingestion.models import (
    DocumentCandidate,
    DocumentPage,
    ExchangeFormat,
    ExtractedRecordSet,
    ExtractedRow,
    IngestionLimits,
    RejectedFileError,
)

def extract_tabular(path: Path, exchange_format: ExchangeFormat, limits: IngestionLimits) -> ExtractedRecordSet:
    rows: list[ExtractedRow] = []
    sheets: list[str] = []
    sheet_headers: list[tuple[str, tuple[object, ...]]] = []
    headers: tuple[object, ...] = ()
    cells = 0

    def append_row(sheet: str, number: int, values: tuple[object, ...]) -> None:
        nonlocal cells
        if len(rows) >= limits.max_tabular_rows or len(values) > limits.max_tabular_columns:
            raise RejectedFileError("TABULAR_LIMIT_EXCEEDED")
        cells += len(values)
        if cells > limits.max_tabular_cells:
            raise RejectedFileError("TABULAR_LIMIT_EXCEEDED")
        rows.append(ExtractedRow(sheet, number, values))

    if exchange_format is ExchangeFormat.CSV:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            sample = source.read(limits.csv_sample_bytes)
            if not sample:
                return ExtractedRecordSet((), (), ("CSV",), (("CSV", ()),))
            try:
                delimiter = determine_delimiter(sample)
            except CsvStructureError as exc:
                raise RejectedFileError(exc.cause_code) from exc
            source.seek(0)
            reader = csv.reader(source, delimiter=delimiter or ",", strict=True)
            headers = tuple(next(reader, ()))
            if len(headers) > limits.max_tabular_columns:
                raise RejectedFileError("TABULAR_LIMIT_EXCEEDED")
            sheet_headers.append(("CSV", headers))
            for number, values in enumerate(reader, start=2):
                append_row("CSV", number, tuple(values))
        sheets.append("CSV")
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                sheets.append(sheet.title)
                iterator = sheet.iter_rows(values_only=True)
                first = tuple(next(iterator, ()))
                if len(first) > limits.max_tabular_columns:
                    raise RejectedFileError("TABULAR_LIMIT_EXCEEDED")
                if not headers:
                    headers = first
                sheet_headers.append((sheet.title, first))
                for number, values in enumerate(iterator, start=2):
                    append_row(sheet.title, number, tuple(values))
        finally:
            workbook.close()
    return ExtractedRecordSet(headers, tuple(rows), tuple(sheets), tuple(sheet_headers))


def extract_document(path: Path, exchange_format: ExchangeFormat) -> DocumentCandidate:
    if exchange_format is ExchangeFormat.PDF:
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            raise RejectedFileError("PROTECTED_DOCUMENT")
        if not reader.pages:
            raise RejectedFileError("EMPTY_DOCUMENT")
        pages: list[DocumentPage] = []
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            pages.append(DocumentPage(number, text, not bool(text)))
        native = "\n\n".join(page.native_text for page in pages if page.native_text)
        state = "PENDING_OCR" if any(page.requires_ocr for page in pages) else "NATIVE_TEXT"
        return DocumentCandidate(tuple(pages), native, state)

    document = Document(path)
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = "\t".join(cell.text.strip() for cell in row.cells)
            if value.strip():
                blocks.append(value)
    native = "\n".join(blocks)
    if not native:
        raise RejectedFileError("EMPTY_DOCUMENT")
    return DocumentCandidate((DocumentPage(1, native, False),), native, "NATIVE_TEXT")
