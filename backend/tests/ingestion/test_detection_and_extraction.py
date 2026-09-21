from __future__ import annotations

from pathlib import Path
from io import BytesIO
import zipfile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
import pytest

from app.ingestion.detection import detect_format
from app.ingestion.extraction import extract_document, extract_tabular
from app.ingestion.models import CONTROLLED_LOCATIONS, ExchangeFormat, IngestionLimits, RejectedFileError, SourceFamily
from app.ingestion.service import IngestionService


@pytest.fixture
def limits() -> IngestionLimits:
    return IngestionLimits(52_428_800, 2_000, 268_435_456, 67_108_864, 100, 65_536, 1_048_576, 250_000, 256, 5_000_000)


def test_csv_detection_is_deterministic_and_preserves_rows(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "litigation.csv"
    source.write_text("id;amount\nL-1;20\nL-2;\n", encoding="utf-8-sig")
    detection = detect_format(source, source.name, limits)
    records = extract_tabular(source, ExchangeFormat.CSV, limits)
    assert detection.accepted and detection.exchange_format is ExchangeFormat.CSV
    assert records.headers == ("id", "amount")
    assert [row.values for row in records.rows] == [("L-1", "20"), ("L-2", "")]


@pytest.mark.parametrize(
    ("name", "content", "cause"),
    [
        ("payload.exe", b"MZ-binary", "UNSUPPORTED_EXTENSION"),
        ("payload.csv", b"a,b;c\n1,2;3\n", "AMBIGUOUS_DELIMITER"),
        ("payload.csv", b"a\x00b", "BINARY_CONTENT"),
    ],
)
def test_unsupported_or_ambiguous_content_is_rejected_safely(
    tmp_path: Path, limits: IngestionLimits, name: str, content: bytes, cause: str,
) -> None:
    source = tmp_path / name
    source.write_bytes(content)
    result = detect_format(source, name, limits)
    assert not result.accepted
    assert result.safe_cause_code == cause


def test_utf8_single_column_csv_is_accepted_and_extracted(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "single.csv"
    source.write_bytes("\ufeffdescription\nfirst value\nsecond value\n".encode("utf-8"))

    detection = detect_format(source, source.name, limits)
    records = extract_tabular(source, ExchangeFormat.CSV, limits)

    assert detection.accepted
    assert records.headers == ("description",)
    assert [row.values for row in records.rows] == [("first value",), ("second value",)]


def test_csv_with_inconsistent_structural_delimiter_is_rejected(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "inconsistent.csv"
    source.write_text("id,value\nonly-one-field\n", encoding="utf-8")

    detection = detect_format(source, source.name, limits)

    assert not detection.accepted
    assert detection.safe_cause_code == "UNDETERMINABLE_STRUCTURE"


def test_controlled_location_does_not_follow_a_file_symlink(tmp_path: Path, limits: IngestionLimits, monkeypatch) -> None:
    root = tmp_path / "controlled"
    location = root / "litigation"
    location.mkdir(parents=True)
    external = tmp_path / "outside.csv"
    external.write_text("id,value\n1,external\n", encoding="utf-8")
    link = location / "link.csv"
    try:
        link.symlink_to(external)
    except OSError:
        pytest.skip("El host no permite crear enlaces simbólicos para esta verificación")
    service = IngestionService(None, None, root / "objects", limits)  # type: ignore[arg-type]
    received: list[bytes] = []

    def receive(stream: BytesIO, **kwargs: object) -> object:
        received.append(stream.read())
        return object()

    monkeypatch.setattr(service, "ingest_stream", receive)
    assert service.run_location(root, "litigation", None) == []  # type: ignore[arg-type]
    assert received == []


def test_xlsx_detection_and_read_only_extraction_preserve_sheet_order(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "compliance.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Obligations"
    first.append(["id", "deadline"])
    first.append(["O-1", "2026-10-01"])
    second = workbook.create_sheet("Evidence")
    second.append(["id", "value"])
    second.append(["O-1", None])
    workbook.save(source)
    assert detect_format(source, source.name, limits).exchange_format is ExchangeFormat.XLSX
    records = extract_tabular(source, ExchangeFormat.XLSX, limits)
    assert records.sheets == ("Obligations", "Evidence")
    assert dict(records.sheet_headers) == {
        "Obligations": ("id", "deadline"),
        "Evidence": ("id", "value"),
    }
    assert [(row.sheet, row.row_number) for row in records.rows] == [("Obligations", 2), ("Evidence", 2)]


def test_pdf_without_native_text_is_a_document_candidate_pending_ocr(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with source.open("wb") as target:
        writer.write(target)
    assert detect_format(source, source.name, limits).exchange_format is ExchangeFormat.PDF
    candidate = extract_document(source, ExchangeFormat.PDF)
    assert candidate.processing_state == "PENDING_OCR"
    assert candidate.pages[0].requires_ocr


def test_zero_page_pdf_and_empty_docx_are_empty_documents(tmp_path: Path) -> None:
    pdf = tmp_path / "empty.pdf"
    with pdf.open("wb") as target:
        PdfWriter().write(target)
    with pytest.raises(RejectedFileError, match="archivo") as pdf_error:
        extract_document(pdf, ExchangeFormat.PDF)
    assert pdf_error.value.cause_code == "EMPTY_DOCUMENT"

    docx = tmp_path / "empty.docx"
    Document().save(docx)
    with pytest.raises(RejectedFileError) as docx_error:
        extract_document(docx, ExchangeFormat.DOCX)
    assert docx_error.value.cause_code == "EMPTY_DOCUMENT"


def test_extension_and_container_mismatch_never_reaches_normal_parser(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "renamed.pdf"
    workbook = Workbook()
    workbook.save(source)
    result = detect_format(source, source.name, limits)
    assert result.exchange_format is ExchangeFormat.XLSX
    assert not result.accepted
    assert result.safe_cause_code == "FORMAT_MISMATCH"


def test_corrupt_supported_container_preserves_its_contractual_format(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "broken.pdf"
    source.write_bytes(b"%PDF-1.7\nnot-a-valid-pdf")
    result = detect_format(source, source.name, limits)
    assert result.exchange_format is ExchangeFormat.PDF
    assert result.supported
    assert not result.accepted
    assert result.safe_cause_code == "CORRUPT_PDF"


def test_controlled_locations_map_only_to_the_four_source_families() -> None:
    assert CONTROLLED_LOCATIONS == {
        "contracts-documents": SourceFamily.CONTRACTS_DOCUMENTS,
        "litigation": SourceFamily.LITIGATION,
        "compliance": SourceFamily.COMPLIANCE,
        "internal-audit": SourceFamily.INTERNAL_AUDIT,
    }


def test_header_only_and_empty_csv_preserve_empty_record_semantics(tmp_path: Path, limits: IngestionLimits) -> None:
    header_only = tmp_path / "header.csv"
    header_only.write_text("id,name\n", encoding="utf-8")
    records = extract_tabular(header_only, ExchangeFormat.CSV, limits)
    assert records.headers == ("id", "name")
    assert records.rows == ()

    empty = tmp_path / "empty.csv"
    empty.write_bytes(b"")
    assert detect_format(empty, empty.name, limits).accepted
    assert extract_tabular(empty, ExchangeFormat.CSV, limits).rows == ()


def test_archive_compression_limit_is_checked_before_extraction(tmp_path: Path, limits: IngestionLimits) -> None:
    source = tmp_path / "oversized.docx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"x")
        archive.writestr("word/document.xml", b"0" * 100_000)
    strict_limits = IngestionLimits(
        limits.max_file_bytes,
        limits.max_archive_entries,
        limits.max_archive_uncompressed_bytes,
        limits.max_archive_entry_bytes,
        2,
        limits.csv_sample_bytes,
        limits.stream_chunk_bytes,
        limits.max_tabular_rows,
        limits.max_tabular_columns,
        limits.max_tabular_cells,
    )
    result = detect_format(source, source.name, strict_limits)
    assert not result.accepted
    assert result.safe_cause_code == "ARCHIVE_COMPRESSION_RATIO_EXCEEDED"


def test_docx_native_text_is_emitted_without_ocr(tmp_path: Path) -> None:
    source = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("Cláusula de vigencia")
    document.save(source)
    candidate = extract_document(source, ExchangeFormat.DOCX)
    assert candidate.processing_state == "NATIVE_TEXT"
    assert candidate.native_text == "Cláusula de vigencia"
    assert not candidate.pages[0].requires_ocr


def test_mixed_pdf_preserves_native_text_and_marks_only_blank_pages_for_ocr(monkeypatch, tmp_path: Path) -> None:
    class Page:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        is_encrypted = False
        pages = [Page("Texto nativo"), Page("")]

    monkeypatch.setattr("app.ingestion.extraction.PdfReader", lambda *args, **kwargs: Reader())
    candidate = extract_document(tmp_path / "mixed.pdf", ExchangeFormat.PDF)
    assert candidate.native_text == "Texto nativo"
    assert [page.requires_ocr for page in candidate.pages] == [False, True]
    assert candidate.processing_state == "PENDING_OCR"
