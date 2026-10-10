"""Evidencia HTTP y de servicio para tablas sin filas, sin conexión PostgreSQL."""
from __future__ import annotations

from contextlib import contextmanager
import hashlib
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from app.ingestion.api import ingestion_service_for, router
from app.ingestion.models import IngestionLimits, SourceFamily
from app.ingestion.service import IngestionService
from app.security.api import current_principal
from app.security.models import AuthenticatedPrincipal
from app.validation.service import TabularRecord, ValidationContext, ValidationService


_HEADERS = b"ID_Contrato,Fecha_Solicitud,Fecha_Vencimiento,Estado_Revision\n"
_CASES = (
    ("zero_bytes", b"", "RECHAZADO", 0),
    ("nonempty_no_header", b"C-EMPTY,2026-10-01,2026-11-01,No iniciado\n", "RECHAZADO", 0),
    ("empty_header", b",,,\n", "RECHAZADO", 0),
    ("incomplete_header", b"ID_Contrato,Fecha_Solicitud\n", "RECHAZADO", 0),
    ("valid_header_only", _HEADERS, "COMPLETADO", 0),
    ("valid_rows", _HEADERS + b"C-VALID,2026-10-01,2026-11-01,No iniciado\n", "COMPLETADO", 1),
)


class _Repository:
    """Retiene metadatos para probar la frontera real sin simular SQL."""

    def __init__(self) -> None:
        self.files = []
        self.rows = []
        self.quarantine = []

    @contextmanager
    def transaction(self):
        yield object()

    def lock_source(self, *_args):
        pass

    def find_existing(self, _connection, **kwargs):
        for received in self.files:
            if received["source_locator"] == kwargs["source_locator"] and received["staged"].sha256 == kwargs["content_sha256"]:
                return {
                    "id": received["file_id"], "operation_id": received["operation_id"],
                    "correlation_id": received["correlation_id"], "state": received["state"],
                    "exchange_format": received["detection"].exchange_format.value,
                    "source_family": received["family"], "safe_cause_code": received["safe_cause_code"],
                }, False
        return None, False

    def next_revision(self, *_args):
        return len(self.files) + 1

    def insert_file(self, _connection, **kwargs):
        self.files.append(kwargs)

    def insert_rows(self, _connection, _file_id, records):
        self.rows.extend(records.rows)

    def insert_quarantine(self, _connection, **kwargs):
        self.quarantine.append(kwargs)
        return uuid4()


class _Security:
    def __init__(self) -> None:
        self.repository = self
        self.events = []

    def require_functional_permission(self, principal, capability):
        assert capability == "ingest.upload"
        return principal

    def revalidate_functional_access(self, _connection, principal, capability):
        assert capability == "ingest.upload"
        return principal

    def write_audit_event(self, _connection, **kwargs):
        self.events.append(kwargs)


def _client(tmp_path: Path):
    repository = _Repository()
    security = _Security()
    principal = AuthenticatedPrincipal(
        uuid4(), uuid4(), "analyst", 1, frozenset({"ANALISTA"}), frozenset({"ingest.upload"}),
    )
    limits = IngestionLimits(52_428_800, 2_000, 268_435_456, 67_108_864, 100, 65_536, 1_048_576, 250_000, 256, 5_000_000)
    service = IngestionService(repository, security, tmp_path / "storage", limits)
    application = FastAPI()
    application.state.security_service = security
    application.include_router(router)
    application.dependency_overrides[current_principal] = lambda: principal
    application.dependency_overrides[ingestion_service_for] = lambda: service
    return TestClient(application), repository, security


def _upload(client, content: bytes, location: str = "contracts-documents", name: str = "contracts.csv"):
    return client.post(
        "/api/ingestion/uploads", files={"file": (name, content)},
        data={"controlled_location": location}, headers={"Idempotency-Key": "same-upload"},
    )


@pytest.mark.parametrize("_name,content,state,row_count", _CASES, ids=[case[0] for case in _CASES])
def test_empty_tabular_structure_is_checked_before_ingestion_completion(tmp_path, _name, content, state, row_count):
    # SRS_REQUIRED: RF-008 rechaza estructura inválida; no prohíbe una tabla válida sin filas.
    client, repository, security = _client(tmp_path)
    response = _upload(client, content)
    assert response.status_code == 201
    body = response.json()
    assert body["state"] == state
    assert body["routing_target"] == ("NONE" if state == "RECHAZADO" else "VALIDATION")
    assert body["safe_cause_code"] == ("STRUCTURAL_INCONSISTENCY" if state == "RECHAZADO" else None)
    assert len(repository.files) == 1
    assert len(repository.rows) == row_count
    assert repository.quarantine == []
    received = repository.files[0]
    assert received["technical_result"] == ("REJECTED" if state == "RECHAZADO" else "ACCEPTED")
    assert received["staged"].path.read_bytes() == content
    assert received["staged"].sha256 == hashlib.sha256(content).digest()
    assert [(event["action"], event["result"]) for event in security.events] == [
        ("INGESTION_COMPLETED", "REJECTED" if state == "RECHAZADO" else "ACCEPTED"),
    ]
    event = security.events[0]
    assert event["resource_identifier"] == body["file_id"]
    assert str(event["correlation_id"]) == body["correlation_id"]
    assert event["safe_cause_code"] == body["safe_cause_code"]


@pytest.mark.parametrize("_name,content,state,_row_count", _CASES, ids=[case[0] for case in _CASES])
def test_empty_tabular_retry_preserves_file_result_and_audit_once(tmp_path, _name, content, state, _row_count):
    # ROBUSTNESS: el rechazo estructural conserva la idempotencia de recepción existente.
    client, repository, security = _client(tmp_path)
    first = _upload(client, content).json()
    counts = (len(repository.files), len(repository.rows), len(security.events))
    repeat = _upload(client, content).json()
    assert first["state"] == repeat["state"] == state
    assert first["idempotent"] is False
    assert repeat["idempotent"] is True
    for key in ("file_id", "operation_id", "correlation_id", "safe_cause_code", "routing_target"):
        assert first[key] == repeat[key]
    assert (len(repository.files), len(repository.rows), len(security.events)) == counts
    assert repository.files[0]["staged"].path.read_bytes() == content
    assert len(list((tmp_path / "storage" / "objects").rglob("*.csv"))) == 1


@pytest.mark.parametrize("headers,state", [((), "RECHAZADO"), (("ID_Contrato", "Fecha_Solicitud", "Fecha_Vencimiento", "Estado_Revision"), "COMPLETADO")])
def test_xlsx_without_rows_uses_same_family_structure_rule(tmp_path, headers, state):
    workbook = Workbook()
    if headers:
        workbook.active.append(headers)
    stream = BytesIO()
    workbook.save(stream)
    client, repository, _security = _client(tmp_path)
    response = _upload(client, stream.getvalue(), name="contracts.xlsx")
    assert response.status_code == 201
    assert response.json()["state"] == state
    assert repository.rows == []


def test_ingestion_with_rows_keeps_downstream_family_validation_boundary(tmp_path):
    # La extracción con filas genéricas sigue delegando validación de familia al downstream.
    client, repository, _security = _client(tmp_path)
    response = _upload(client, b"id,amount\nL-1,20\n", location="litigation", name="litigation.csv")
    assert response.status_code == 201
    assert response.json()["state"] == "COMPLETADO"
    assert response.json()["routing_target"] == "VALIDATION"
    assert len(repository.rows) == 1

    class DownstreamRepository(_Repository):
        def family_for_file(self, _connection, _file_id):
            return SourceFamily.LITIGATION

        def write_audit_event(self, _connection, **kwargs):
            self.events.append(kwargs)

    downstream = DownstreamRepository()
    downstream.events = []
    security = _security
    result = ValidationService(downstream, security).validate(
        file_id=uuid4(), family=SourceFamily.LITIGATION, headers=("id", "amount"),
        records=(TabularRecord(2, uuid4(), {"id": "L-1", "amount": "20"}, 2),),
        context=ValidationContext(uuid4(), uuid4(), security.events[0]["actor"]),
        capability="ingest.upload",
    )
    assert result.file_rejected
    assert result.validated_records == ()
    assert result.conforming_positions == ()
    assert downstream.rows == []
    assert downstream.quarantine == []
    assert [(event["action"], event["result"]) for event in downstream.events] == [("VALIDATION_FILE_REJECTED", "REJECTED")]


@pytest.mark.parametrize("location,headers", [
    ("contracts-documents", _HEADERS),
    ("litigation", b"ID_Litigio,Fecha_Apertura,Estado,Nivel_Severidad\n"),
    ("compliance", b"ID_Obligacion,Fecha_Limite\n"),
    ("internal-audit", b"ID_Incidente,Fecha_Evento,Area,Nivel_Severidad\n"),
    ("internal-audit", b"ID_Asunto,Tipo_Asunto,Estado,Fecha\n"),
])
def test_header_only_is_accepted_for_each_existing_family_contract(tmp_path, location, headers):
    client, repository, _security = _client(tmp_path)
    response = _upload(client, headers, location=location)
    assert response.status_code == 201
    assert response.json()["state"] == "COMPLETADO"
    assert response.json()["safe_cause_code"] is None
    assert repository.rows == []


@pytest.mark.parametrize("content", [b"\xef\xbb\xbf", b"\n"])
def test_bom_only_or_blank_header_is_not_a_valid_empty_table(tmp_path, content):
    client, repository, _security = _client(tmp_path)
    response = _upload(client, content)
    assert response.status_code == 201
    assert response.json()["state"] == "RECHAZADO"
    assert response.json()["safe_cause_code"] == "STRUCTURAL_INCONSISTENCY"
    assert repository.rows == []


def test_empty_xlsx_validates_every_sheet_header(tmp_path):
    workbook = Workbook()
    workbook.active.append(("ID_Contrato", "Fecha_Solicitud", "Fecha_Vencimiento", "Estado_Revision"))
    workbook.create_sheet("Incomplete").append(("ID_Contrato",))
    stream = BytesIO()
    workbook.save(stream)
    client, repository, _security = _client(tmp_path)
    response = _upload(client, stream.getvalue(), name="contracts.xlsx")
    assert response.status_code == 201
    assert response.json()["state"] == "RECHAZADO"
    assert response.json()["routing_target"] == "NONE"
    assert repository.rows == []
