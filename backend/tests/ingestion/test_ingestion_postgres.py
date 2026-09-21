"""Evidencia PostgreSQL de atomicidad, idempotencia y concurrencia."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
import base64
from datetime import UTC, datetime, timedelta
from io import BytesIO
import os
from pathlib import Path
from uuid import uuid4

from alembic import command
from alembic.config import Config
import pytest
import sqlalchemy as sa
from sqlalchemy.engine import make_url
from openpyxl import Workbook
from docx import Document
from pypdf import PdfWriter

from app.ingestion.models import IdempotencyConflictError, IngestionLimits, RoutingTarget
from app.ingestion.repository import IngestionRepository
from app.ingestion.service import IngestionService
from app.security.models import AuthenticatedPrincipal, AuthorizationError
from app.security.repository import SecurityRepository
from app.security.service import SecurityService
from app.security.tokens import JwtService


BACKEND_ROOT = Path(__file__).resolve().parents[2]


def _url() -> str:
    value = os.getenv("DATA_TEST_DATABASE_URL", "")
    if not value:
        pytest.skip("DATA_TEST_DATABASE_URL no configurada")
    if "test" not in (make_url(value).database or "").lower():
        pytest.fail("La evidencia de ingesta requiere una base desechable")
    return value


@pytest.fixture
def database(tmp_path: Path):
    url = _url()
    parsed = make_url(url)
    os.environ.update(
        POSTGRES_HOST=parsed.host or "localhost", POSTGRES_PORT=str(parsed.port or 5432),
        POSTGRES_USER=parsed.username or "", POSTGRES_PASSWORD=parsed.password or "",
        POSTGRES_DB=parsed.database or "",
    )
    config = Config(str(BACKEND_ROOT / "alembic.ini"))
    command.downgrade(config, "base")
    command.upgrade(config, "head")
    engine = sa.create_engine(url)
    account_id, session_id = uuid4(), uuid4()
    with engine.begin() as connection:
        connection.execute(sa.text("INSERT INTO app.user_account (id, username, display_name, password_hash) VALUES (:id, 'analyst', 'Analyst', 'hash')"), {"id": account_id})
        connection.execute(sa.text("INSERT INTO app.user_role (user_id, role_id) VALUES (:id, 'ANALISTA')"), {"id": account_id})
        connection.execute(
            sa.text("""INSERT INTO app.access_session
                (id, user_id, refresh_token_sha256, authorization_version, state, expires_at)
                VALUES (:session, :user, :hash, 1, 'ACTIVE', :expires)"""),
            {"session": session_id, "user": account_id, "hash": bytes(32), "expires": datetime.now(UTC) + timedelta(hours=1)},
        )
    principal = AuthenticatedPrincipal(account_id, session_id, "analyst", 1, frozenset({"ANALISTA"}), frozenset({"ingest.upload", "ingest.execute"}))
    conninfo = parsed.set(drivername="postgresql").render_as_string(hide_password=False)
    encoded_secret = base64.urlsafe_b64encode(bytes(range(32))).decode("ascii")
    security = SecurityService(
        SecurityRepository(conninfo),
        JwtService(issuer="issuer", audience="audience", keyring={"key": encoded_secret}, active_kid="key"),
    )
    limits = IngestionLimits(52_428_800, 2_000, 268_435_456, 67_108_864, 100, 65_536, 1_048_576, 250_000, 256, 5_000_000)
    service = IngestionService(IngestionRepository(conninfo), security, tmp_path / "objects", limits)
    try:
        yield engine, service, principal
    finally:
        engine.dispose()
        command.downgrade(config, "base")


@pytest.mark.requires_db
@pytest.mark.contract
def test_manual_idempotency_and_new_source_revision_are_enforced(database) -> None:
    engine, service, principal = database
    first = service.ingest_stream(
        BytesIO(b"id,amount\nL-1,20\n"), original_name="litigation.csv", controlled_location="litigation",
        principal=principal, capability="ingest.upload", source_locator="manual/analyst/key-1", idempotency_key="key-1",
    )
    repeated = service.ingest_stream(
        BytesIO(b"id,amount\nL-1,20\n"), original_name="litigation.csv", controlled_location="litigation",
        principal=principal, capability="ingest.upload", source_locator="manual/analyst/key-1", idempotency_key="key-1",
    )
    assert repeated.idempotent and repeated.file_id == first.file_id
    with pytest.raises(IdempotencyConflictError):
        service.ingest_stream(
            BytesIO(b"id,amount\nL-1,30\n"), original_name="litigation.csv", controlled_location="litigation",
            principal=principal, capability="ingest.upload", source_locator="manual/analyst/key-1", idempotency_key="key-1",
        )

    service.ingest_stream(
        BytesIO(b"id,amount\nL-1,20\n"), original_name="litigation.csv", controlled_location="litigation",
        principal=principal, capability="ingest.execute", source_locator="controlled/litigation/current.csv", idempotency_key=None,
    )
    service.ingest_stream(
        BytesIO(b"id,amount\nL-1,30\n"), original_name="litigation.csv", controlled_location="litigation",
        principal=principal, capability="ingest.execute", source_locator="controlled/litigation/current.csv", idempotency_key=None,
    )
    with engine.connect() as connection:
        assert connection.execute(sa.text("SELECT count(*) FROM app.ingest_file WHERE id = :id"), {"id": first.file_id}).scalar_one() == 1
        assert connection.execute(sa.text("SELECT array_agg(source_revision ORDER BY source_revision) FROM app.ingest_file WHERE source_locator = 'controlled/litigation/current.csv'")).scalar_one() == [1, 2]


@pytest.mark.requires_db
@pytest.mark.contract
def test_concurrent_detection_creates_one_effective_operation(database) -> None:
    engine, service, principal = database

    def ingest():
        return service.ingest_stream(
            BytesIO(b"id,deadline\nO-1,2026-10-01\n"), original_name="compliance.csv", controlled_location="compliance",
            principal=principal, capability="ingest.execute", source_locator="controlled/compliance/daily.csv", idempotency_key=None,
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: ingest(), range(2)))
    assert sum(result.idempotent for result in results) == 1
    with engine.connect() as connection:
        assert connection.execute(sa.text("SELECT count(*) FROM app.ingest_file WHERE source_locator = 'controlled/compliance/daily.csv'")).scalar_one() == 1


@pytest.mark.requires_db
@pytest.mark.contract
def test_audit_failure_rolls_back_file_and_database_state(database, monkeypatch) -> None:
    engine, service, principal = database

    def fail_audit(*args, **kwargs):
        raise RuntimeError("controlled audit failure")

    monkeypatch.setattr(service.security.repository, "write_audit_event", fail_audit)
    with pytest.raises(RuntimeError, match="controlled audit failure"):
        service.ingest_stream(
            BytesIO(b"id,amount\nL-1,20\n"), original_name="litigation.csv", controlled_location="litigation",
            principal=principal, capability="ingest.upload", source_locator="manual/analyst/rollback", idempotency_key="rollback",
        )
    with engine.connect() as connection:
        assert connection.execute(sa.text("SELECT count(*) FROM app.ingest_file WHERE source_locator = 'manual/analyst/rollback'")).scalar_one() == 0
        assert connection.execute(sa.text("SELECT count(*) FROM audit.event WHERE action = 'INGESTION_COMPLETED'")).scalar_one() == 0


@pytest.mark.requires_db
@pytest.mark.contract
def test_multisheet_workbook_preserves_sheet_headers_rows_and_order(database) -> None:
    engine, service, principal = database
    workbook = Workbook()
    obligations = workbook.active
    obligations.title = "Obligations"
    obligations.append(["id", "deadline"])
    obligations.append(["O-1", "2026-10-01"])
    evidence = workbook.create_sheet("Evidence")
    evidence.append(["id", "value"])
    evidence.append(["O-1", "present"])
    content = BytesIO()
    workbook.save(content)
    content.seek(0)

    result = service.ingest_stream(
        content,
        original_name="compliance.xlsx",
        controlled_location="compliance",
        principal=principal,
        capability="ingest.upload",
        source_locator="manual/analyst/workbook",
        idempotency_key="workbook",
    )
    assert result.routing_target is RoutingTarget.VALIDATION
    with engine.connect() as connection:
        rows = connection.execute(
            sa.text(
                "SELECT source_sheet, row_number, raw_payload "
                "FROM app.source_record WHERE ingest_file_id = :file ORDER BY source_sheet, row_number"
            ),
            {"file": result.file_id},
        ).all()
    assert [(row.source_sheet, row.row_number) for row in rows] == [("Evidence", 2), ("Obligations", 2)]
    assert rows[0].raw_payload == {"id": "O-1", "value": "present"}
    assert rows[1].raw_payload == {"deadline": "2026-10-01", "id": "O-1"}


@pytest.mark.requires_db
@pytest.mark.contract
def test_corrupt_document_enters_quarantine_with_minimized_audit(database) -> None:
    engine, service, principal = database
    result = service.ingest_stream(
        BytesIO(b"%PDF-1.7\ninvalid"),
        original_name="contract.pdf",
        controlled_location="contracts-documents",
        principal=principal,
        capability="ingest.upload",
        source_locator="manual/analyst/corrupt-document",
        idempotency_key="corrupt-document",
    )
    assert result.state == "CUARENTENA"
    assert result.routing_target is RoutingTarget.VALIDATION
    assert result.safe_cause_code == "CORRUPT_PDF"
    with engine.connect() as connection:
        quarantine = connection.execute(
            sa.text(
                "SELECT q.cause_code, q.original_object_id, f.stored_object_id "
                "FROM app.quarantine_item q JOIN app.ingest_file f ON f.id = q.ingest_file_id "
                "WHERE q.ingest_file_id = :file"
            ),
            {"file": result.file_id},
        ).one()
        events = connection.execute(
            sa.text(
                "SELECT action, result, safe_cause_code FROM audit.event "
                "WHERE correlation_id = :correlation ORDER BY occurred_at, action"
            ),
            {"correlation": result.correlation_id},
        ).all()
    assert quarantine.cause_code == "CORRUPT_PDF"
    assert quarantine.original_object_id == quarantine.stored_object_id
    assert {(event.action, event.result, event.safe_cause_code) for event in events} == {
        ("INGESTION_COMPLETED", "FAILED", "CORRUPT_PDF"),
        ("QUARANTINE_CREATED", "SUCCESS", "CORRUPT_PDF"),
    }


@pytest.mark.requires_db
@pytest.mark.contract
@pytest.mark.parametrize("name", ["empty.pdf", "empty.docx"])
def test_empty_documents_enter_quarantine_with_a_stable_cause(database, name: str) -> None:
    engine, service, principal = database
    if name.endswith(".pdf"):
        content = BytesIO()
        PdfWriter().write(content)
    else:
        content = BytesIO()
        Document().save(content)
    content.seek(0)

    result = service.ingest_stream(
        content, original_name=name, controlled_location="contracts-documents", principal=principal,
        capability="ingest.upload", source_locator=f"manual/analyst/{name}", idempotency_key=name,
    )

    assert result.state == "CUARENTENA"
    assert result.safe_cause_code == "EMPTY_DOCUMENT"
    assert result.routing_target is RoutingTarget.VALIDATION
    assert result.document is None
    with engine.connect() as connection:
        row = connection.execute(
            sa.text("SELECT technical_result, safe_cause_code FROM app.ingest_file WHERE id = :id"),
            {"id": result.file_id},
        ).one()
        quarantine = connection.execute(
            sa.text("SELECT cause_code FROM app.quarantine_item WHERE ingest_file_id = :id"),
            {"id": result.file_id},
        ).one()
        events = connection.execute(
            sa.text("SELECT action FROM audit.event WHERE correlation_id = :id"),
            {"id": result.correlation_id},
        ).scalars().all()
    assert row == ("FAILED", "EMPTY_DOCUMENT")
    assert quarantine.cause_code == "EMPTY_DOCUMENT"
    assert set(events) == {"QUARANTINE_CREATED", "INGESTION_COMPLETED"}


@pytest.mark.requires_db
@pytest.mark.contract
def test_resource_limit_rejection_is_audited_without_storing_a_partial_object(database, monkeypatch) -> None:
    engine, service, principal = database
    service.limits = IngestionLimits(6, 2_000, 268_435_456, 67_108_864, 100, 65_536, 1_048_576, 250_000, 256, 5_000_000)
    exact = service.ingest_stream(
        BytesIO(b"value\n"), original_name="exact.csv", controlled_location="litigation", principal=principal,
        capability="ingest.upload", source_locator="manual/analyst/exact", idempotency_key="exact",
    )
    assert exact.state == "COMPLETADO"
    monkeypatch.setattr("app.ingestion.service.detect_format", lambda *args: pytest.fail("parser invoked"))

    result = service.ingest_stream(
        BytesIO(b"value\nX"), original_name="large.csv", controlled_location="litigation", principal=principal,
        capability="ingest.upload", source_locator="manual/analyst/large", idempotency_key="large",
    )

    assert result.state == "RECHAZADO"
    assert result.safe_cause_code == "RESOURCE_LIMIT_EXCEEDED"
    with engine.connect() as connection:
        row = connection.execute(
            sa.text("""SELECT stored_object_id, content_sha256, content_sha256_complete, observed_byte_size,
                              format_classification, technical_result, safe_cause_code
                       FROM app.ingest_file WHERE id = :id"""),
            {"id": result.file_id},
        ).one()
        event = connection.execute(
            sa.text("SELECT result, safe_cause_code FROM audit.event WHERE correlation_id = :id"),
            {"id": result.correlation_id},
        ).one()
    assert row == (None, None, False, 7, "UNDETERMINED", "REJECTED", "RESOURCE_LIMIT_EXCEEDED")
    assert event == ("REJECTED", "RESOURCE_LIMIT_EXCEEDED")


@pytest.mark.requires_db
@pytest.mark.contract
def test_resource_limit_rejection_rolls_back_when_its_audit_event_fails(database, monkeypatch) -> None:
    engine, service, principal = database
    service.limits = IngestionLimits(5, 2_000, 268_435_456, 67_108_864, 100, 65_536, 1_048_576, 250_000, 256, 5_000_000)

    def fail_audit(*args, **kwargs):
        raise RuntimeError("controlled audit failure")

    monkeypatch.setattr(service.security.repository, "write_audit_event", fail_audit)
    with pytest.raises(RuntimeError, match="controlled audit failure"):
        service.ingest_stream(
            BytesIO(b"abcdeX"), original_name="large.csv", controlled_location="litigation", principal=principal,
            capability="ingest.upload", source_locator="manual/analyst/large-rollback", idempotency_key="large-rollback",
        )
    with engine.connect() as connection:
        assert connection.execute(
            sa.text("SELECT count(*) FROM app.ingest_file WHERE source_locator = 'manual/analyst/large-rollback'")
        ).scalar_one() == 0


@pytest.mark.requires_db
@pytest.mark.contract
def test_permission_revocation_is_revalidated_and_denial_is_audited_without_mutation(database) -> None:
    engine, service, principal = database
    with engine.begin() as connection:
        connection.execute(
            sa.text(
                "DELETE FROM app.role_permission "
                "WHERE role_id = 'ANALISTA' AND permission_id = 'ingest.upload'"
            )
        )

    with pytest.raises(AuthorizationError):
        service.ingest_stream(
            BytesIO(b"id,amount\nL-1,20\n"),
            original_name="litigation.csv",
            controlled_location="litigation",
            principal=principal,
            capability="ingest.upload",
            source_locator="manual/analyst/revoked",
            idempotency_key="revoked",
        )

    with engine.connect() as connection:
        assert connection.execute(
            sa.text("SELECT count(*) FROM app.ingest_file WHERE source_locator = 'manual/analyst/revoked'")
        ).scalar_one() == 0
        event = connection.execute(
            sa.text(
                "SELECT action, result, resource_identifier FROM audit.event "
                "WHERE action = 'INGESTION_DENIED'"
            )
        ).one()
    assert event.action == "INGESTION_DENIED"
    assert event.result == "DENIED"
    assert event.resource_identifier is None
